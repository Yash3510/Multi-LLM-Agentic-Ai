"""
Exercise every 4CE tool against the real runtime.

    cd backend && ../.venv/bin/python ../4ce/test_tools.py           # macOS, Linux
    cd backend && ../.venv/Scripts/python.exe ../4ce/test_tools.py   # Windows

Run it from the backend directory so the open_webui package and its database
resolve. Each tool is loaded from source the same way the application loads it,
then driven through its success and failure paths. Prints a pass/fail summary
and exits non-zero if anything fails, so it can gate a demo-day check.
"""

import asyncio
import glob
import importlib.util
import json
import os
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path

# Running this script puts its own directory on sys.path, not the working
# directory, so the backend package would not resolve without this.
sys.path.insert(0, os.getcwd())

# Tools record what they write in 4CE's audit trail. A test run is not work
# the plant did, so it records to a trail of its own, not data/4ce/audit.jsonl.
os.environ["FOURCE_AUDIT_PATH"] = str(Path(tempfile.gettempdir()) / "4ce-test-audit.jsonl")

TOOLS = Path(__file__).resolve().parent / "tools"
RESULTS: list[tuple[str, bool, str]] = []


def load(name: str):
    spec = importlib.util.spec_from_file_location(name, TOOLS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.Tools()


def check(label: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((label, ok, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))


def admin_user_id() -> str | None:
    matches = glob.glob("data/webui.db")
    if not matches:
        return None
    row = sqlite3.connect(matches[0]).execute("select id from user limit 1").fetchone()
    return row[0] if row else None


async def test_sandbox() -> None:
    print("\n4CE Sovereign Sandbox")
    tool = load("sandbox")

    out = await tool.run_python("print('hello'); print(6*7)")
    check("executes code and returns stdout", "42" in out and "completed" in out)

    out = await tool.run_python(
        "import urllib.request\n"
        "try:\n"
        "    urllib.request.urlopen('https://example.com', timeout=5)\n"
        "    print('REACHABLE')\n"
        "except Exception as exc:\n"
        "    print('blocked:', type(exc).__name__)"
    )
    check("outbound network is blocked", "REACHABLE" not in out and "blocked:" in out)

    out = await tool.run_python("raise ValueError('boom')")
    check("non-zero exit reported as failed", "failed" in out and "ValueError" in out)

    out = await tool.run_python("open('/workspace/x','w')")
    check("workspace mount is read-only", "failed" in out)

    original = tool.valves.timeout_seconds
    tool.valves.timeout_seconds = 5
    out = await tool.run_python("import time; time.sleep(30)")
    check("wall-clock timeout enforced", "Timed out" in out or "time limit" in out)
    tool.valves.timeout_seconds = original

    out = await tool.run_python("   ")
    check("empty source rejected", "No code" in out)

    # Files are how a run hands back a spreadsheet or a chart. /output is the only
    # writable path that survives, and a configured directory keeps a copy on disk.
    with tempfile.TemporaryDirectory(prefix="4ce-test-out-") as keep:
        tool.valves.output_dir = keep
        out = await tool.run_python(
            "open('/output/readings.csv','w').write('a,b'); print('done')"
        )
        check(
            "files written to /output are returned",
            "completed" in out and "Files produced" in out and "readings.csv" in out,
        )
        landed = Path(keep) / "readings.csv"
        check("returned files reach the configured directory", landed.is_file())
        check(
            "the returned file holds what the run wrote",
            landed.is_file() and landed.read_text() == "a,b",
        )

    out = await tool.run_python("print('no files here')")
    check("a run producing no files says so", "Files produced" not in out)

    original_cap = tool.valves.max_output_file_mb
    tool.valves.max_output_file_mb = 0
    out = await tool.run_python("open('/output/big.bin','wb').write(b'x'*2048); print('ok')")
    check("an oversized file is refused, not silently dropped", "Not returned" in out)
    tool.valves.max_output_file_mb = original_cap


async def test_deliverables() -> None:
    print("\n4CE Deliverables")
    tool = load("deliverables")
    uid = admin_user_id()
    if not uid:
        check("admin user available", False, "no user row found; run from backend/")
        return

    out = await tool.create_word_document(
        "Approval Note - Pump P-101B",
        "## Scope\n\nRoutine inspection.\n\n## Findings\n\n- Seal weeping 3-4 dpm\n- Vibration Zone B\n\n"
        "## Recommendation\n\nApprove return to service.",
        "MIR/2026/0431",
        __user__={"id": uid},
    )
    ok = "Download" in out and ".docx" in out
    check("generates a .docx and returns a link", ok, out.splitlines()[0][:70] if not ok else "")

    if ok:
        newest = max(glob.glob("data/uploads/*.docx"), key=lambda p: Path(p).stat().st_mtime)
        with zipfile.ZipFile(newest) as archive:
            names = archive.namelist()
            body = archive.read("word/document.xml").decode("utf-8", "replace")
        check("file is a valid OOXML package", "word/document.xml" in names)
        check("content is present in the document", "P-101B" in body and "Recommendation" in body)
        check("classification banner applied", "CONFIDENTIAL" in body)

    out = await tool.create_word_document("Title", "   ", __user__={"id": uid})
    check("empty body rejected", "empty" in out.lower())

    # An image the answer shows is embedded, not left as a link nobody
    # reading the file away from 4CE could open.
    import io
    import uuid as _uuid
    from docx import Document
    from PIL import Image

    swatch = io.BytesIO()
    Image.new("RGB", (400, 200), (250, 248, 240)).save(swatch, "PNG")
    image_id = str(_uuid.uuid4())
    body = "Findings.\n\n![Each field boxed where it was read](/api/v1/files/" + image_id + "/content)\n"
    spec = importlib.util.spec_from_file_location("deliverables_module", TOOLS / "deliverables.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    blocks = module._blocks(body)
    check("an image line is a block of its own", [b[0] for b in blocks] == ["para", "image"])
    plain = Document(io.BytesIO(tool._build_docx("T", body, "", images={})))
    shown = Document(io.BytesIO(tool._build_docx("T", body, "", images={image_id: swatch.getvalue()})))
    check("the Word report embeds the image", len(shown.inline_shapes) == len(plain.inline_shapes) + 1)
    deck = module._build_pptx("T", "### Read from the image\n\n" + body, "Deliverable", [], [],
                              {image_id: swatch.getvalue()})
    from pptx import Presentation
    pictures = [s for slide in Presentation(io.BytesIO(deck)).slides for s in slide.shapes if s.shape_type == 13]
    check("the deck gives the image a slide", len(pictures) == 1)

    out = await tool.create_word_document("", "content", __user__={"id": uid})
    check("missing title rejected", "title is required" in out.lower())

    out = await tool.create_word_document("Title", "content", __user__=None)
    check("unattributed request rejected", "not generated" in out.lower())


async def test_sovereignty() -> None:
    print("\n4CE Sovereignty Check")
    tool = load("sovereignty")

    out = await tool.verify_sovereignty()
    check("produces an audit table", "Sovereignty audit" in out and "| Surface |" in out)
    check("reports a verdict", "PASS" in out or "ATTENTION REQUIRED" in out)
    # Nothing is watching egress in this process yet: the report must say so,
    # not print a verdict that sounds observed.
    check("says when egress was not observed", "PASS on configuration, not observed" in out
          or "ATTENTION REQUIRED" in out)
    check("states its own limitation", "rests on configuration alone" in out)

    from open_webui.utils.fource_egress import watch
    watch.set_endpoints(["http://127.0.0.1:1234/v1"])
    watch.start()
    await asyncio.sleep(0.8)
    try:
        out = await tool.verify_sovereignty()
        check("reports what the workbench was observed doing",
              "What the workbench actually did" in out and "samples" in out)
        check("states what sampling cannot see", "not by capturing packets" in out)
    finally:
        watch.stop()

    ok_local = tool._is_local("http://localhost:1234/v1") and tool._is_local("http://127.0.0.1:8080")
    check("classifies loopback as on-premise", ok_local)
    check("classifies private ranges as on-premise", tool._is_local("http://192.168.1.50:11434"))
    check("classifies public hosts as external", not tool._is_local("https://api.openai.com/v1"))


DEMO_REPORT = """1.2  Mechanical seal exhibits intermittent weeping, approximately 3 to 4 drops
     per minute during sustained operation.
1.3  Bearing housing temperature recorded at 71 degrees C against an alarm limit of 80 degrees C.
1.4  Vibration measured at 4.1 mm/s RMS. ISO 10816 Zone B.
1.5  Coupling guard fastener missing at position 3 of 4.
Minimum measured wall thickness: 11.2 mm (nominal 12.7 mm)
Retirement thickness: 9.5 mm
"""


async def test_egress() -> None:
    """The egress watch's classification, against a socket table it is handed.

    Synthetic, so it is deterministic and makes no connection of its own: the
    point is that each kind of flow lands in the right count, once.
    """
    print("\n4CE Egress watch")
    from collections import namedtuple

    import psutil
    import open_webui.utils.fource_egress as egress

    conn = namedtuple("conn", "fd family type laddr raddr status pid")
    addr = namedtuple("addr", "ip port")
    me, model = os.getpid(), 999_001

    def c(local, remote, status="ESTABLISHED", pid=me):
        return conn(-1, 2, 1, addr(*local), addr(*remote) if remote else (), status, pid)

    table: list = []
    listening = [
        c(("127.0.0.1", 8080), None, psutil.CONN_LISTEN),
        c(("127.0.0.1", 1234), None, psutil.CONN_LISTEN, model),
    ]
    real_table, real_family = egress.psutil.net_connections, egress._family
    egress.psutil.net_connections = lambda kind="inet": listening + table
    egress._family = lambda pid, excluded=None: {pid: ("backend.exe" if pid == me else "model.exe", "")} \
        if pid in (me, model) else {}
    try:
        w = egress.EgressWatch()
        w.set_endpoints(["http://127.0.0.1:1234/v1", "http://192.168.1.50:8000/v1"])

        def sample(*flows):
            table[:] = list(flows)
            w._sample()

        sample(c(("127.0.0.1", 50001), ("127.0.0.1", 1234)))
        sample(c(("127.0.0.1", 50001), ("127.0.0.1", 1234)))
        check("a loopback flow is counted once across samples", w.flows["local"] == 1)
        check("the model server is in scope by its port",
              {m["role"] for m in w._scope.values()} == {"backend", "model server"})

        sample(c(("192.168.1.9", 8080), ("192.168.1.20", 61000)))
        sample(c(("192.168.1.9", 50100), ("192.168.1.50", 8000)))
        check("LAN clients and the configured LAN endpoint are not leaks",
              w.flows["lan_in"] == 1 and w.flows["lan_expected"] == 1 and w.flows["lan_out"] == 0)

        mark = w.mark()
        sample(c(("192.168.1.9", 50200), ("192.168.1.77", 3128)))
        sample(c(("10.0.0.5", 50300), ("93.184.216.34", 443), pid=model))
        event = next(iter(w._events["external"].values()))
        check("other outbound LAN and the internet are flagged",
              w.flows["lan_out"] == 1 and w.flows["external"] == 1)
        check("an external flow names the process that made it", event["role"] == "model server")
        check("a run is not reported as observed when nothing was watching",
              w.since_mark(mark)["observed"] is False)

        w._thread = type("Alive", (), {"is_alive": lambda self: True})()
        run = w.since_mark(mark)
        check("a watched run reports what happened during it",
              run["observed"] and run["external"] == 1 and run["lan_out"] == 1)

        import time as _time
        w._canary = ("1.1.1.1", 443, _time.time() + 5)
        sample(c(("10.0.0.5", 50400), ("1.1.1.1", 443)))
        check("the canary is logged apart, not counted external",
              w.flows["canary"] == 1 and w.flows["external"] == 1)

        class Proc:
            def __init__(self, pid, name, kids=()):
                self.pid, self._name, self._kids = pid, name, list(kids)

            def name(self):
                return self._name

            def children(self):
                return self._kids

        left_out: dict = {}
        tree = Proc(1, "Bionic.exe", [Proc(2, "llmster.exe", [Proc(3, "node.exe")]),
                                      Proc(4, "brave.exe", [Proc(5, "brave.exe")])])
        kept = [p.pid for p in egress._descendants(tree, left_out)]
        check("a browser the model server opened is not the model server",
              kept == [2, 3] and left_out == {4: "brave.exe"})

        w.reset()
        check("a new window starts from zero", w.samples == 0 and w.flows["external"] == 0)
        check("the snapshot is serialisable for the page", bool(json.dumps(w.snapshot())))
    finally:
        egress.psutil.net_connections, egress._family = real_table, real_family


async def test_audit_trail() -> None:
    """The audit trail: append-only, each entry sealed to the one before."""
    print("\n4CE Audit trail")
    import hashlib

    import open_webui.utils.fource_audit as audit

    with tempfile.TemporaryDirectory(prefix="4ce-test-trail-") as folder:
        path = Path(folder) / "audit.jsonl"
        trail = audit.AuditTrail(path)
        token = audit._run.set({"chat": "chat-1", "message": "msg-1"})
        try:
            first = trail.append("request", task="document", prompt=hashlib.sha256(b"q").hexdigest())
            trail.append("model.call", agent="FRIDAY", model="qwen/qwen3-vl-4b", seconds=1.5)
            third = trail.append("release", fingerprint="ab" * 32)
        finally:
            audit._run.reset(token)
        check("the first entry follows the genesis hash", first["seq"] == 1 and first["prev"] == audit.GENESIS)
        check("each entry names the one before it", third["prev"] == trail.entries()[1]["hash"])
        check("entries carry the run they belong to", all(e["chat"] == "chat-1" for e in trail.entries()))
        result = trail.verify()
        check("an untouched trail verifies", result["ok"] and result["entries"] == 3 and result["head"] == third["hash"])

        again = audit.AuditTrail(path)
        fourth = again.append("approval", decision="approved")
        check("a reopened trail carries on the same chain", fourth["seq"] == 4 and fourth["prev"] == third["hash"])
        check("entries after a number", [e["seq"] for e in again.entries(after=2)] == [3, 4])

        lines = path.read_text(encoding="utf-8").splitlines()
        changed = json.loads(lines[1])
        changed["model"] = "some/other-model"
        path.write_text("\n".join([lines[0], json.dumps(changed), *lines[2:]]) + "\n", encoding="utf-8")
        result = again.verify()
        check("an entry changed after writing breaks the chain there",
              not result["ok"] and result["broken_at"] == 2 and "changed" in result["reason"])

        path.write_text("\n".join([lines[0], lines[1], lines[3]]) + "\n", encoding="utf-8")
        result = again.verify()
        check("a removed entry breaks the chain there", not result["ok"] and result["broken_at"] == 3)

        blocked = audit.AuditTrail(Path(folder))  # a directory: nothing can be appended to it
        check("a trail that cannot be written reports it, and does not raise",
              blocked.append("request") is None and blocked.failed_writes == 1 and blocked.last_error)

    # Offline mode: a library may not install a model while it runs. Proved on
    # a stand-in module, then on the loader that actually did it.
    import open_webui.utils.fource_offline as offline

    with tempfile.TemporaryDirectory(prefix="4ce-test-guard-") as folder:
        (Path(folder) / "fource_fake_loader.py").write_text("def fetch():\n    return 'downloaded'\n")
        sys.path.insert(0, folder)
        offline.GUARDED["fource_fake_loader"] = "fetch"
        try:
            offline.guard()
            import fource_fake_loader

            try:
                fource_fake_loader.fetch()
                refused = False
            except RuntimeError as exc:
                refused = "offline mode" in str(exc)
            check("a guarded download is refused when its module is first imported", refused)
        finally:
            offline.GUARDED.pop("fource_fake_loader", None)
            sys.path.remove(folder)
            sys.modules.pop("fource_fake_loader", None)
    import unstructured.nlp.tokenize as tokenize

    check("the document loader's spaCy download is refused offline",
          getattr(tokenize._install_spacy_model, "__fource_guard__", False))

    spec = importlib.util.spec_from_file_location(
        "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
    )
    orchestrator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(orchestrator)
    check("model calls are recorded under the agent that made them",
          orchestrator._agent_of(orchestrator._FRIDAY_SYSTEM) == "FRIDAY"
          and orchestrator._agent_of(orchestrator._READER_SYSTEM) == "READER")
    check("the trail's hashes are SHA-256 of the text",
          orchestrator._digest("abc") == hashlib.sha256(b"abc").hexdigest() == audit.digest("abc"))


async def test_workspace_tools() -> None:
    """File read and write inside one workspace; spreadsheets read with their
    formulas and changed only in a copy."""
    print("\n4CE Workspace files and spreadsheets")
    import shutil

    files, sheets = load("files"), load("sheets")
    with tempfile.TemporaryDirectory(prefix="4ce-test-ws-") as folder:
        files.valves.workspace = sheets.valves.workspace = folder

        out = await files.write_file("notes/p101b.md", "# P-101B\nfirst")
        check("writes a text file inside the workspace", out.startswith("**Written**")
              and (Path(folder) / "notes" / "p101b.md").read_text() == "# P-101B\nfirst")
        out = await files.write_file("notes/p101b.md", "# P-101B\nsecond")
        kept = Path(folder) / ".versions" / "notes" / "p101b.md" / "p101b.v1.md"
        check("overwriting keeps the earlier version", "earlier version kept" in out
              and kept.read_text() == "# P-101B\nfirst")
        refused = [await files.write_file(p, "x") for p in ("../escape.md", "C:/Windows/x.md", "/etc/x.md")]
        check("a path outside the workspace is refused", all("outside the workspace" in r for r in refused)
              and not (Path(folder).parent / "escape.md").exists())
        check("an executable type is refused", "not a type 4CE writes" in await files.write_file("run.bat", "x"))
        check("the kept versions cannot be rewritten",
              "kept, not rewritten" in await files.write_file(".versions/notes/p101b.md/p101b.v1.md", "x"))
        out = await files.read_file("notes/p101b.md")
        check("reads a file back", "second" in out and "from the workspace" in out)
        check("a read outside the workspace is refused", "outside" in await files.read_file("../../x.md"))

        survey = Path(__file__).parent / "demo" / "samples" / "thickness_survey_P-101B.xlsx"
        shutil.copy(survey, Path(folder) / "survey.xlsx")
        text = await sheets.read_sheet("survey.xlsx")
        check("a workbook is read with its cell letters and headers",
              "B · Previous thickness (mm)" in text and "| Shell course 1, east | 12 | 11.2 | 9.5 | 3 | 2 |" in text)

        spec = importlib.util.spec_from_file_location(
            "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
        )
        o = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(o)
        changes = o._survey_changes(text)
        check("a survey's formulas use its own columns",
              changes and changes[0]["formula"] == '=IF(E{row}>0,(B{row}-C{row})/E{row},"")'
              and "/F{row}" in changes[1]["formula"])
        out = await sheets.write_sheet(json.dumps(changes), "survey.xlsx")
        check("changes go to a copy, as live formulas", out.startswith("**Excel workbook**")
              and (Path(folder) / "survey - 4CE.xlsx").exists() and "was not changed" in out)
        check("the source is untouched", (Path(folder) / "survey.xlsx").read_bytes() == survey.read_bytes())
        check("the copy's formulas give what 4CE's calculation gives",
              "| Shell course 1, east | 0.266667 | 6.375 | 3.1875 | 2 |" in out
              and "| Bottom head | 0.6 | 0.666667 | 0.333333 | 6 |" in out)

        # Each location through the rule pack: its own margin, its own interval.
        rows = o._survey_rows(text)
        check("a survey's locations are read with their sheet rows",
              rows[0] == (2, "Shell course 1, east", 11.2, 9.5) and len(rows) == 5)
        sop = load("sop_check")
        intervals = {}
        for row, _, cur, req in rows:
            assessed = await sop.check_sop_thresholds(
                f"Minimum measured wall thickness {cur:g} mm. Retirement thickness {req:g} mm."
            )
            years, basis = o._sop_interval(assessed)
            if years:
                intervals[row] = years
        check("§4.2's six months falls on each thin location, not the sheet as a whole",
              intervals == {2: 0.5, 4: 0.5, 5: 0.5, 6: 0.5})
        worked = await load("calculations").calculate_remaining_life(
            o._with_interval_column(text, intervals), 0.0, basis
        )
        check("each location's next measurement honours its own interval",
              "| Discharge nozzle N2 | 10 | 9.9 | 8.2 | 3 | 0.033 | 51.0 | 0.5 |" in worked
              and "| Shell course 1, west | 12 | 11.6 | 9.5 | 3 | 0.133 | 15.7 | 7.9 |" in worked)
        before = sorted(p.name for p in Path(folder).iterdir())
        preview = await sheets.write_sheet(json.dumps(o._survey_changes(text, intervals)), "survey.xlsx",
                                           __preview__=True)
        check("a preview shows what approval would write, and writes nothing",
              preview.startswith("On approval") and "| Discharge nozzle N2 | 0.0333333 | 51 | 0.5 | 0.5 | 5 |" in preview
              and sorted(p.name for p in Path(folder).iterdir()) == before)
        planned = o._survey_changes(text, intervals)
        claims = o._workbook_checks(
            "The values are already present in the spreadsheet; no new columns are required.", planned
        )
        check("an answer saying the workbook already has the results is a problem",
              len(claims) == 1 and claims[0]["kind"] == "problem" and "adds 4 columns" in claims[0]["text"])
        said = ("The 3-year interval is authoritative and not subject to revision per SOP-MEC-014 §4.2, "
                "which only applies to margin < 2.0 mm — a condition not met in any location.")
        found = o._interval_checks(said, intervals, "the interval SOP-MEC-014 §4.2 requires", 5)
        check("an answer saying §4.2 applies nowhere, where it applies at four locations, is a problem",
              len(found) == 1 and "4 of 5 locations" in found[0]["text"])
        check("an answer that states §4.2 as it applies is left alone", o._interval_checks(
            "SOP-MEC-014 §4.2 requires a six-month interval where the margin is under 2.0 mm.",
            intervals, "the interval SOP-MEC-014 §4.2 requires", 5) == [])
        check("an answer that reports the results is left alone",
              o._workbook_checks("Shell course 1, east corrodes at 0.267 mm/year; 6.4 years remain.", planned) == [])
        check("4CE's own sections are not weighed as claims",
              o._answer_only("Answer.\n\n### The workbook, on approval\n\n| H · SOP interval (yr) |") == "Answer.\n\n")
        out = await sheets.write_sheet(json.dumps(o._survey_changes(text, intervals)), "survey.xlsx",
                                       save_as="survey with intervals.xlsx")
        check("the copy carries the interval and the next inspection as formulas",
              "| Discharge nozzle N2 | 0.0333333 | 51 | 0.5 | 0.5 | 5 |" in out
              and "| Shell course 1, west | 0.133333 | 15.75 |  | 7.875 | 3 |" in out)
        copy_text = await sheets.read_sheet("survey - 4CE.xlsx")
        check("the copy reads back with its formulas", "(=IF(E2>0,(B2-C2)/E2,\"\"))" in copy_text)
        calc = load("calculations")
        worked = await calc.calculate_remaining_life(text)
        check("a sheet read by the tool feeds the calculation", "| Shell course 1, east | 12 | 11.2 | 9.5 | 3 | 0.267 | 6.4 |" in worked)

        out = await sheets.write_sheet(json.dumps([{"cell": "Z1", "formula": '=WEBSERVICE("http://x")'}]), "survey.xlsx")
        check("a formula that could reach outside the workbook is refused", "refused" in out and "Nothing was written" in out)
        out = await sheets.write_sheet(json.dumps(changes), "survey.xlsx", save_as="survey.xlsx")
        check("the source is never overwritten", "never changes" in out)
        (Path(folder) / "log.csv").write_text("Tag,Reading\nP-101A,18.6\nP-101B,17.9\n")
        out = await sheets.write_sheet(json.dumps([{"column": "Double", "formula": "=B{row}*2"}]), "log.csv")
        check("a CSV is changed as a workbook copy", "copied into a workbook" in out and "| P-101B | 35.8 | 3 |" in out)

    check("an edit request is told apart from a question",
          o._wants_sheet_edit("Add the corrosion rate and remaining life to this survey")
          and not o._wants_sheet_edit("What is the remaining life at the bottom head?"))
    check("changes JARVIS proposes are read from its block",
          o._sheet_block('Done.\n```4ce-sheet\n[{"cell": "G2", "value": 1}]\n```') == [{"cell": "G2", "value": 1}])
    check("an attached workbook is found by name", o._attached_sheets(
        {"files": [{"type": "file", "name": "survey.xlsx"}, {"type": "file", "name": "photo.png"}]}) == ["survey.xlsx"])


async def test_execution_check() -> None:
    """The orchestrator's own check on generated code, from real sandbox runs."""
    print("\n4CE Orchestrator: execution check")
    spec = importlib.util.spec_from_file_location(
        "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
    )
    orchestrator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(orchestrator)
    judge = orchestrator._execution_check
    sandbox = load("sandbox")

    code = "def double(x):\n    return 2 * x\nassert double(4) == 8\nprint(double(4))"
    ran = judge(code, await sandbox.run_python(code))
    check("a clean run with its assertions holding passes", ran["kind"] == "ok" and "1 assertion" in ran["text"])

    code = "assert 2 + 2 == 5, 'arithmetic'"
    ran = judge(code, await sandbox.run_python(code))
    check("a failed assertion fails, with its message", ran["kind"] == "problem" and "arithmetic" in ran["text"])

    code = "print(undefined_name)"
    ran = judge(code, await sandbox.run_python(code))
    check("a crash fails, naming the error", ran["kind"] == "problem" and "NameError" in ran["text"])

    ran = judge("print(1)", "The sovereign sandbox is unavailable, so this code was NOT executed.\n\n"
                            "Reason: the Docker daemon is not running\n")
    check("unexecuted code is unverified, and final", ran["kind"] == "unverified" and ran.get("final"))
    check("no code block is a problem", judge("", "")["kind"] == "problem")

    code = "def median(xs):\n    return sorted(xs)[len(xs) // 2]\nmedian([5, 3, 9, 1, 7])"
    ran = judge(code, await sandbox.run_python(code), "Write a function and print it for [5, 3, 9, 1, 7].")
    check("asked to print, printed nothing, is a problem", ran["kind"] == "problem" and "printed nothing" in ran["text"])

    code = "print(sorted([5, 3, 9, 1, 7])[2])"
    ran = judge(code, await sandbox.run_python(code), "Print the median of [5, 3, 9, 1, 7].")
    check("a clean run that asserts nothing is not a pass", ran["kind"] == "unverified")


async def test_figure_check() -> None:
    """The orchestrator's own check that cited figures are in what they cite."""
    print("\n4CE Orchestrator: figure check")
    spec = importlib.util.spec_from_file_location(
        "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
    )
    orchestrator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(orchestrator)
    judge = orchestrator._figure_checks
    sources = [{"name": "SOP-MEC-014_seal_leakage.txt", "passages": [{"text":
        "2.2 Leakage between 5 and 20 drops per minute: replace within 30 days."}]}]

    kinds = lambda checks: sorted((c["kind"], c["text"].split(" ")[0]) for c in checks)
    checks = judge("Replace within 30 days [1].", sources)
    check("a figure the source holds passes", kinds(checks) == [("ok", "Every")])
    checks = judge("Replace within 14 days [1].", sources)
    check("a figure the source lacks fails", ("problem", "14") in kinds(checks))
    checks = judge("The seal leaks 18 drops per minute [1].", sources, "it leaks 18 drops/min")
    check("the request's own figure is shown, not failed",
          ("unverified", "18") in kinds(checks) and not any(k == "problem" for k, _ in kinds(checks)))
    checks = judge("Clause 5.1 forbids starting without the guard [1].", sources, "", {"5.1"})
    check("a clause the rule pack cited is not a fabrication", not any(c["kind"] == "problem" for c in checks))
    checks = judge("Clause 9.4 forbids it [1].", sources, "", {"5.1"})
    check("a clause nobody cited still fails", any("Clause 9.4" in c["text"] for c in checks))

    checks = judge("SOP-MEC-014 limits seal leakage to 7 drops per minute.", sources)
    check("no citation, no claim: an invented SOP limit fails", any(
        c["kind"] == "problem" and c["text"].startswith("7 is stated as the SOP's") for c in checks))
    checks = judge("SOP-MEC-014 requires replacement within 30 days.", sources)
    check("an uncited SOP figure a source holds is not called invented",
          not any(c["kind"] == "problem" for c in checks))

    # A shift-log answer: a clock time, a figure given the wrong document, and
    # two documents that share a code.
    both = sources + [{"name": "SOP-MEC-014_readings_P-101B.txt", "passages": [{"text":
        "Seal weeping 3-4 drops per minute. Coupling guard fastener missing."}]}]
    checks = judge("The guard was found at 06:05 missing a fastener [2].", both)
    check("a clock time is not a figure", not any(c["text"].startswith("06") for c in checks))
    checks = judge("Schedule the replacement within 30 days [2].", both)
    check("a figure in another source is shown as mis-cited, not invented",
          any(c["kind"] == "unverified" and c["text"] == "30 is in SOP-MEC-014 · seal leakage, but is cited to "
              "SOP-MEC-014 · readings P-101B" for c in checks) and not any(c["kind"] == "problem" for c in checks))
    ruled = "| Seal leakage | 6 drops/min | below 5 drops/min | §2.2 | REVIEW |\n- **§2.2** - Replace within 45 days."
    checks = judge("Replace within 45 days [1].", sources, "", set(), "", ruled)
    check("a rule-pack figure is not called invented",
          any("the SOP rule pack's figure" in c["text"] for c in checks) and not any(c["kind"] == "problem" for c in checks))
    checks = judge("Replace within 30 days [2]. Then within 30 days [2].", both)
    check("one mis-citation is listed once", sum(c["text"].startswith("30 ") for c in checks) == 1)
    check("a document code shared by two sources is followed by the rest of the name",
          orchestrator._doc_label("SOP-MEC-014_readings_P-101B.txt", [s["name"] for s in both])
          == "SOP-MEC-014 · readings P-101B"
          and orchestrator._doc_label("SOP-MEC-014_seal_leakage.txt") == "SOP-MEC-014")

    table = ("| Parameter | Measured | Acceptance limit | Clause | Verdict |\n|---|---|---|---|---|\n"
             "| Seal leakage | 6 drops/min | below 5 drops/min | §2.2 | REVIEW |\n"
             "| Coupling guard | not reported | no coupling guard fastener missing | §5.1 | PASS |")
    verdicts = orchestrator._verdict_checks
    found = verdicts("Seal leakage at 6 drops per minute is within acceptable limits under Clause 2.1.", table)
    check("an answer calling a REVIEW reading within limits is a problem",
          len(found) == 1 and found[0]["kind"] == "problem" and found[0]["on"] == "verdict" and "§2.2" in found[0]["text"])
    check("an answer that agrees with the rule pack is left alone",
          verdicts("Seal leakage at 6 drops per minute exceeds the acceptable limit of 5.", table) == []
          and verdicts("The coupling guard is within limits.", table) == [])

    relevant = orchestrator._relevant
    passages = [{"text": "Seal leakage below 5 drops per minute is acceptable.", "name": "SOP-MEC-014", "distance": 0.68},
                {"text": "4. Wall thickness", "name": "SOP-MEC-014", "distance": 0.86}]
    check("an unrelated question keeps only strong matches",
          [p["distance"] for p in relevant(passages, "What is the capital of France?", 0.8)] == [0.86])
    check("a related question keeps what it shares words with",
          len(relevant(passages, "What is the acceptable seal leakage?", 0.8)) == 2)

    sop = ("## SOP assessment - SOP-MEC-014 - ATTENTION REQUIRED\n### Required actions\n"
           "- **§4.2** - Remaining wall thickness margin is below 2.0 mm. Shorten the inspection interval to six months.")
    years, basis = orchestrator._sop_interval(sop)
    check("reads the interval an SOP requires", abs(years - 0.5) < 1e-9 and basis == "the interval SOP-MEC-014 §4.2 requires")
    check("no required interval, none read",
          orchestrator._sop_interval("## SOP assessment - X\n- **§2.2** - Replace the seal.") == (0.0, ""))


async def test_calculations() -> None:
    """Remaining life by the thickness method, worked in code."""
    print("\n4CE Engineering Calculations")
    tool = load("calculations")

    out = await tool.calculate_remaining_life(
        "previous thickness 12.0 mm three years ago, current thickness 11.2 mm, required thickness 9.5 mm"
    )
    check("reproduces the worked example: 0.267 mm/year, 6.4 years, 3.2 years",
          "**0.267 mm/year**" in out and "**6.4 years**" in out and "**3.2 years**" in out)
    check("shows every step with its units", "(12.0 mm - 11.2 mm) ÷ 3 years" in out and "Step 3" in out)

    out = await tool.calculate_remaining_life(
        "readings 3 years apart\n\n| TML | Previous (mm) | Current (mm) | Required (mm) |\n|---|---|---|---|\n"
        "| TML-1 | 12.0 | 11.2 | 9.5 |\n| TML-3 | 10.4 | 9.4 | 9.5 |"
    )
    check("a survey table is worked row by row", "| TML-1 |" in out and "| TML-3 |" in out)
    check("at or below required thickness governs, and says so",
          "Governing location: TML-3" in out and "AT OR BELOW REQUIRED THICKNESS" in out)

    out = await tool.calculate_remaining_life(
        "previous thickness 12.0 mm 3 years ago, current thickness 11.2 mm, required thickness 9.5 mm, "
        "maximum interval 2 years"
    )
    check("the code maximum caps the next measurement",
          "The lesser of RL ÷ 2 (3.2 years) and the code maximum (2.0 years) = **2.0 years**" in out)

    out = await tool.calculate_remaining_life(
        "previous thickness 12.0 mm 3 years ago, current thickness 11.2 mm, required thickness 9.5 mm",
        max_interval_years=0.5, interval_basis="the interval SOP-MEC-014 §4.2 requires",
    )
    check("an SOP's shorter interval governs, and is named",
          "the interval SOP-MEC-014 §4.2 requires (6 months) = **6 months**" in out
          and "again within 6 months" in out)

    out = await tool.calculate_remaining_life("the pump leaks 18 drops/min")
    check("says what it needs when readings are missing", out.startswith("Remaining life could not be calculated"))

    import io
    from openpyxl import load_workbook
    spec = importlib.util.spec_from_file_location("deliverables_module", TOOLS / "deliverables.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    book = load_workbook(io.BytesIO(module._build_xlsx(
        "T", "x", [], [], [{"kind": "remaining_life", "max_interval": None, "rows": [
            {"location": "TML-1", "previous": 12.0, "current": 11.2, "required": 9.5, "years": 3.0}]}]
    )))
    sheet = book["Remaining life"]
    check("the workbook carries it as live formulas",
          str(sheet["G2"].value).startswith("=") and str(sheet["H2"].value).startswith("=IF("))


async def test_routing() -> None:
    """The router reads the model registry, not code."""
    print("\n4CE Orchestrator: model registry")
    from types import SimpleNamespace

    spec = importlib.util.spec_from_file_location(
        "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
    )
    orchestrator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(orchestrator)
    registry = json.loads((Path(__file__).parent / "models.json").read_text(encoding="utf-8"))
    check("the built-in registry is the same as models.json",
          orchestrator._DEFAULT_REGISTRY["models"] == registry["models"]
          and orchestrator._DEFAULT_REGISTRY["routing"] == registry["routing"]
          and orchestrator._DEFAULT_REGISTRY["embedding"] == registry["embedding"])

    served = {"qwen/qwen3-vl-4b": {}, "qwen/qwen3-1.7b": {}, "ace_orchestrator.tony": {}}
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(MODELS=served)))
    pipe = orchestrator.Pipe()
    model, why, candidates = pipe._route("code", request, "ace_orchestrator.tony")
    check("code goes to the model the registry says can code",
          model == "qwen/qwen3-1.7b" and "needs coding" in why)
    check("the others are listed with why they were passed over",
          any(c["id"] == "qwen/qwen3-vl-4b" and c["status"] == "lacks coding" for c in candidates))
    model, _, _ = pipe._route("vision", request, "ace_orchestrator.tony")
    check("a scan goes to the model with vision", model == "qwen/qwen3-vl-4b")

    registry["models"].append({"id": "example/new-coder-7b", "capabilities": ["coding"], "licence": "Apache-2.0"})
    pipe.valves.model_registry = json.dumps(registry)
    _, _, candidates = pipe._route("code", request, "ace_orchestrator.tony")
    check("a new registry entry appears with no code change",
          any(c["id"] == "example/new-coder-7b" and c["status"] == "not served" for c in candidates))

    pipe.valves.coding_model = "qwen/qwen3-vl-4b"
    model, why, _ = pipe._route("code", request, "ace_orchestrator.tony")
    check("an override in the valves still wins, and says so", model == "qwen/qwen3-vl-4b" and "override" in why)


async def test_vision() -> None:
    """Reading an image: its kind, its fields, and what is left to the reviewer."""
    print("\n4CE Orchestrator: reading an image")
    import io
    from PIL import Image

    spec = importlib.util.spec_from_file_location(
        "orchestrator", str(Path(__file__).parent / "functions" / "orchestrator.py")
    )
    o = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(o)

    paper = Image.new("RGB", (1654, 1180), (250, 248, 240))
    pump = Image.new("RGB", (1200, 900), (70, 110, 60))
    check("a P&ID is read as a drawing", o._image_kind("Extract the tags from this P&ID excerpt", paper) == "drawing")
    check("a shift log is read as a page", o._image_kind("Here is last night's shift log", pump) == "page")
    check("paper with no word for it is a page", o._image_kind("What does this say?", paper) == "page")
    check("anything else is a photograph", o._image_kind("What is wrong here?", pump) == "photo")
    check("a page is sent large, a photo small",
          o.Pipe().valves.vision_page_edge >= 2200 and o.Pipe().valves.vision_max_edge == 900)
    check("the size an image is sent at", o._sent_size((2200, 1300), 900) == (900, 532))

    # A reply as qwen3-vl-4b gave it for the shift log, with a line repeated
    # and the last object cut off by the token budget.
    reply = """```json
[
  {"tag": "P-101A", "item": "discharge pressure", "value": "18.6 bar g", "unit": "bar g", "time": "03:45", "confidence": "high", "bbox_2d": [97, 368, 605, 418]},
  {"tag": "P-101B", "item": "seal weeping", "value": "approx 6", "unit": "drops/min", "time": "02:10", "confidence": "medium", "bbox_2d": [97, 173, 703, 224]},
  {"tag": "P-101B", "item": "coupling guard bolt", "value": "missing", "unit": "", "time": "06:05", "confidence": "high", "bbox_2d": [97, 557, 598, 607]},
  {"tag": "P-101A", "item": "discharge pressure", "value": "18.6 bar g", "unit": "bar g", "time": "03:45", "confidence": "high", "bbox_2d": [97, 368, 605, 418]},
  {"tag": "P-101B", "item": "brg temp", "value": "7"""
    fields = o._parse_fields(reply, "page")
    check("every finished field is kept, once, from a reply cut off mid-object", len(fields) == 3)
    check("fields are numbered in reading order", [f["item"] for f in fields][:2] == ["seal weeping", "discharge pressure"])
    check("a unit is written once, in its column", fields[1]["value"] == "18.6" and fields[1]["unit"] == "bar g")
    check("less than full confidence is marked for checking", fields[0]["check"] and not fields[1]["check"])
    table = o._fields_table(fields, "page", paper.size)
    check("the table says where, in the image's pixels", "160,204 → 1163,264" in table and "**check**" in table)
    lines = o._readings_text(fields)
    check("the rule pack gets one line per reading", "02:10 P-101B seal weeping: approx 6 drops/min" in lines)
    sop = load("sop_check")
    assessed = await sop.check_sop_thresholds(lines)
    check("a page's readings reach the rule pack", "§2.2 | REVIEW" in assessed and "§5.1 | FAIL" in assessed)
    checks = o._reading_checks({"fields": fields, "kind": "page"})
    check("the doubtful field is left to the reviewer, by number",
          checks[0]["kind"] == "ok" and any(c["kind"] == "unverified" and "#1" in c["text"] for c in checks))

    drawing = """[
  {"tag": "PT 101", "type": "instrument", "confidence": "high", "bbox_2d": [590, 333, 620, 376]},
  {"tag": "FE-101", "type": "equipment", "confidence": "high", "bbox_2d": [670, 507, 705, 533]},
  {"tag": "TO E-101", "type": "other", "confidence": "high", "bbox_2d": [832, 507, 900, 533]},
  {"tag": "CRUDE FEED TANK", "type": "other", "confidence": "high", "bbox_2d": [57, 549, 151, 576]},
  {"tag": "P&ID EXCERPT - CRUDE CHARGE PUMPS P-101A/B", "type": "other", "confidence": "high", "bbox_2d": [651, 825, 918, 857]},
  {"tag": "HV-101A", "type": "valve", "confidence": "low", "bbox_2d": [1130, 780, 1333, 860]},
  {"tag": "8\\"-P-1001-A1A", "type": "line", "confidence": "high", "bbox_2d": [134, 439, 200, 464]}
]"""
    tags = {f["tag"]: f for f in o._parse_fields(drawing, "drawing", (2200, 1300))}
    check("an ISA bubble read as PT 101 is the tag PT-101", "PT-101" in tags and tags["PT-101"]["type"] == "instrument")
    check("a tag's letters settle its type: FE is an instrument", tags.get("FE-101", {}).get("type") == "instrument")
    check("an off-page connector names its tag", tags.get("E-101", {}).get("type") == "equipment")
    check("labels and the title block are not tags", len(tags) == 5 and "CRUDE FEED TANK" not in tags)
    check("a box in the sent image's pixels comes onto the grid", 500 < tags["HV-101A"]["box"][0] < 530)
    check("a line number is typed as a line", tags.get('8"-P-1001-A1A', {}).get("type") == "line")

    png = o._annotate(paper, fields)
    check("the annotated copy is a PNG the image's size", Image.open(io.BytesIO(png)).size == paper.size)

    # The platform puts a block of file ids in front of a message with a file.
    message = {"role": "user", "content": [
        {"type": "text", "text": '<attached_files>\n<file type="file" id="c5f9849e" name="shift_log_P-101.png"/>\n'
                                 "</attached_files>\n\n"},
        {"type": "text", "text": "Here is last night's handwritten shift log for the P-101 pumps. Read it."},
    ]}
    asked = o._request_text(message)
    check("the request is what the requester wrote, without the file block",
          asked.startswith("Here is last night's") and "attached_files" not in asked)
    check("a hand-over is titled by what is handed over",
          o._document_title(asked) == "Last night's handwritten shift log for the P-101 pumps")
    check("a second request after 'and' is not in the title", o._document_title(
        "Extract every tag and line number from this P&ID excerpt and describe how the pumps connect to E-101."
    ) == "Tag and line number from this P&ID excerpt")


async def test_sop_check() -> None:
    print("\n4CE SOP Threshold Check")
    tool = load("sop_check")

    out = await tool.check_sop_thresholds(DEMO_REPORT)
    check("produces an assessment table", "SOP assessment" in out and "| Parameter |" in out)
    check("seal leakage 3-4 dpm passes on 2.1", "| Seal leakage | 3 to 4 drops/min" in out and "§2.1 | PASS" in out)
    check("vibration Zone B passes on 3.1", "| Zone B |" in out and "§3.1 | PASS" in out)
    check("wall thickness margin computed", "1.7 mm margin" in out)
    check("thin margin raises the 4.2 action", "**§4.2**" in out and "six months" in out)
    check("missing guard fastener fails on 5.1", "§5.1 | FAIL" in out)
    check("blocks the start on a 5.1 failure", "Disposition: NOT FIT FOR START" in out)
    check("unassessed readings are listed", "Readings outside this SOP" in out and "Bearing housing" in out)

    out = await tool.check_sop_thresholds("Seal leakage 12 drops per minute.")
    check("12 dpm requires review under 2.2", "§2.2 | REVIEW" in out and "30 days" in out)

    out = await tool.check_sop_thresholds("Seal leakage 25 drops per minute.")
    check("25 dpm fails under 2.3", "§2.3 | FAIL" in out and "REMOVE FROM SERVICE" in out)

    out = await tool.check_sop_thresholds("Seal leakage 3 to 25 drops per minute.")
    check("a range is assessed at its worse end", "§2.3 | FAIL" in out)

    out = await tool.check_sop_thresholds("Vibration ISO 10816 Zone D.")
    check("Zone D fails under 3.3", "§3.3 | FAIL" in out)

    out = await tool.check_sop_thresholds(
        "Minimum measured wall thickness 9.1 mm. Retirement thickness 9.5 mm."
    )
    check("thickness below retirement fails under 4.1", "§4.1 | FAIL" in out and "-0.4 mm margin" in out)

    out = await tool.check_sop_thresholds("Seal shows atomised spray at the gland.")
    check("atomised spray fails under 2.3", "Seal atomisation" in out and "§2.3 | FAIL" in out)

    out = await tool.check_sop_thresholds("Inspector attended site. Vibration ISO 10816 Zone A.")
    check("an absent reading is NO DATA, not PASS", "| Seal leakage | not found" in out and "NO DATA" in out)
    check("incomplete assessment does not read as fit", "Disposition: ASSESSMENT INCOMPLETE" in out)

    out = await tool.check_sop_thresholds("   ")
    check("empty input rejected", "No readings were supplied" in out)

    broken = load("sop_check")
    broken.valves.rule_pack = "{not json"
    out = await broken.check_sop_thresholds("Seal leakage 3 drops per minute.")
    check("malformed rule pack reported, not ignored", "could not be read" in out and "No assessment" in out)

    custom = load("sop_check")
    custom.valves.rule_pack = json.dumps({
        "sop_id": "SOP-ELE-002",
        "rules": [{
            "id": "ir", "parameter": "Insulation resistance", "unit": "MOhm", "type": "band",
            "limit": "above 50 MOhm",
            "patterns": [r"insulation resistance[^0-9]{0,20}(\d+(?:\.\d+)?)"],
            "bands": [
                {"max": 50, "verdict": "FAIL", "clause": "3.1", "action": "Do not energise.",
                 "disposition": "REMOVE FROM SERVICE"},
                {"verdict": "PASS", "clause": "3.2", "disposition": "FIT FOR SERVICE"},
            ],
        }],
    })
    out = await custom.check_sop_thresholds("Insulation resistance 22 MOhm measured at 500 V.")
    check("a custom rule pack drives the engine", "SOP-ELE-002" in out and "§3.1 | FAIL" in out)

    out = await tool.check_sop_thresholds(DEMO_REPORT)
    check("states its own limitation", "only for the parameters in its rule pack" in out)

    out = await tool.check_sop_thresholds("06:05 Coupling guard bolt missing on P-101B.")
    check("a missing guard bolt, as a shift log puts it, fails on 5.1", "§5.1 | FAIL" in out)
    out = await tool.check_sop_thresholds(
        "Coupling guard fasteners: none missing. No visible spray from the seal. Seal leakage 2 drops/min."
    )
    check("a defect reported absent is not failed", "§5.1 | PASS" in out and "§2.3 | PASS" in out
          and "REMOVE FROM SERVICE" not in out)

    out = await tool.check_sop_thresholds("seal leak 18 drops/min, vibration 7.1 mm/s, bearing temperature 71 C")
    check("a velocity is named, not called missing", "7.1 mm/s was given" in out and "state the zone" in out)
    check("18 drops/min needs the supervisor under §2.2", "| 18 drops/min |" in out and "§2.2 | REVIEW" in out)

    out = await tool.check_sop_thresholds(
        "previous thickness 12.0 mm three years ago, current thickness 11.2 mm, required thickness 9.5 mm"
    )
    check("current and required thickness are read as §4.1's measured and retirement",
          "§4.1 | PASS" in out and "11.2" in out)
    check("a 1.7 mm margin brings §4.2's six-month interval", "§4.2" in out and "six months" in out)


async def main() -> None:
    if not Path("data").exists():
        print("Run this from the backend/ directory so open_webui and its database resolve.")
        sys.exit(2)

    await test_sandbox()
    await test_deliverables()
    await test_sovereignty()
    await test_egress()
    await test_audit_trail()
    await test_workspace_tools()
    await test_execution_check()
    await test_figure_check()
    await test_calculations()
    await test_routing()
    await test_vision()
    await test_sop_check()

    failed = [label for label, ok, _ in RESULTS if not ok]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} checks passed")
    if failed:
        print("Failed: " + "; ".join(failed))
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
